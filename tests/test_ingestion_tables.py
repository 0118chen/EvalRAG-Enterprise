"""Spreadsheet ingestion: a flattened sheet loses which column a number belongs to."""

import datetime
from io import BytesIO

import openpyxl

from app.core.ingestion import chunk_pages, extract_text


def _workbook(sheets: dict[str, list[list[object]]]) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _text(payload: bytes, filename: str = "table.xlsx") -> str:
    pages = extract_text(filename, payload)
    assert len(pages) == 1
    return pages[0][1]


CALCULATION_SHEET = {
    "测算表": [
        ["地区", "中央财政补贴比例", "省级财政补贴比例"],
        ["河南省", 0.4, 0.3],
        ["山东省", 0.35, 0.25],
    ]
}


def _percent_sheet() -> bytes:
    """A sheet whose numbers are formatted as percentages, as Excel stores them."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "测算表"
    for row in CALCULATION_SHEET["测算表"]:
        sheet.append(row)
    for row in (2, 3):
        for column in (2, 3):
            sheet.cell(row=row, column=column).number_format = "0%"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_each_row_repeats_its_column_names() -> None:
    """A row must be readable on its own, wherever the chunker cuts the text."""
    text = _text(_percent_sheet())
    row_line = next(line for line in text.splitlines() if "河南省" in line)

    assert "中央财政补贴比例: 40%" in row_line
    assert "省级财政补贴比例: 30%" in row_line
    assert "地区: 河南省" in row_line


def test_title_rows_above_the_header_are_kept_as_content() -> None:
    """The real sample sheet: 附件1 / 102个县（市）名单 / blank / then the header.

    Those title rows are content, but they are not column names - the first row
    that is more than one cell wide is where the header starts.
    """
    payload = _workbook(
        {
            "附件1": [
                ["附件1"],
                ["102个县（市）名单"],
                [],
                ["序号", "省辖市", "市县区"],
                ["1", "郑州市", "中牟县"],
            ]
        }
    )
    lines = [line for line in _text(payload).splitlines() if line.strip()]

    assert lines[0] == "附件1"
    assert lines[1] == "102个县（市）名单"
    assert "省辖市: 郑州市" in lines[2]
    assert "市县区: 中牟县" in lines[2]
    assert "附件1:" not in lines[2]


def test_sheet_number_becomes_the_page_number() -> None:
    payload = _workbook({"第一张表": [["科目", "金额"], ["利息收入", 120]], "第二张表": [["科目", "余额"], ["应收利息", 30]]})
    pages = extract_text("book.xlsx", payload)

    assert [page for page, _ in pages] == [1, 2]
    assert "利息收入" in pages[0][1]
    assert "应收利息" not in pages[0][1]


def test_empty_cells_are_skipped_instead_of_producing_blank_pairs() -> None:
    text = _text(_workbook({"表": [["地区", "备注", "比例"], ["河南省", None, 0.4]]}))
    row_line = next(line for line in text.splitlines() if "河南省" in line)

    assert "备注:" not in row_line
    assert "备注: " not in row_line
    assert "比例: 0.4" in row_line


def test_percentage_cells_keep_the_displayed_format() -> None:
    """0.4 is stored; 40% is what the document says. The number format carries that."""
    text = _text(_percent_sheet())

    assert "40%" in text
    assert ": 0.4" not in text


def test_a_plain_number_is_not_turned_into_a_percentage() -> None:
    text = _text(_workbook({"表": [["地区", "贴息金额"], ["河南省", 0.4]]}))

    assert "贴息金额: 0.4" in text
    assert "40%" not in text


def test_dates_are_serialised_as_iso_dates() -> None:
    # Excel has no timezone support, so a real workbook only ever holds naive
    # datetimes; openpyxl refuses to even write a tz-aware one.
    naive = datetime.datetime(2024, 3, 1)  # noqa: DTZ001
    payload = _workbook({"表": [["地区", "起息日"], ["河南省", naive]]})
    text = _text(payload)

    assert "起息日: 2024-03-01" in text


def test_every_row_is_retrievable_together_with_its_column_names() -> None:
    """The design claim, stated as a test.

    With a repeated header the window only has to contain the row, not the whole
    sheet; the chunker's overlap is what makes that hold at the boundaries.
    """
    rows: list[list[object]] = [["指标0", "指标1", "指标2", "指标3", "指标4", "指标5"]]
    rows += [[f"地区{index}", index, index + 1, index + 2, index + 3, index + 4] for index in range(40)]
    chunks = chunk_pages("doc-xlsx", extract_text("table.xlsx", _workbook({"测算表": rows})))
    assert len(chunks) > 1

    for index in range(40):
        holding = [chunk.text for chunk in chunks if f"地区{index} " in chunk.text or chunk.text.endswith(f"地区{index}")]
        assert holding, f"row {index} is in no chunk"
        assert any("指标1:" in text for text in holding), f"row {index} lost its column names"


def test_a_single_row_sheet_is_not_silently_empty() -> None:
    """No data rows: the header itself is content and must survive, otherwise the
    document fails the empty-extraction guard for no good reason."""
    text = _text(_workbook({"测算表": [["地区", "中央财政补贴比例", "省级财政补贴比例"]]}))

    assert "测" in text
    assert "中央财政补贴比例" in text
